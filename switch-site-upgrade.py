import requests
import getpass
import json
import logging
import sys
import os
from datetime import datetime, timezone, timedelta
import curses
import openpyxl
import argparse

# Set up logging
logging.basicConfig(
    level=logging.DEBUG,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler("switch_upgrade_debug.log"),
        logging.StreamHandler(sys.stdout)
    ]
)
logger = logging.getLogger(__name__)

logger.debug("Starting script execution.")

upgrade_id_file = "last_upgrade_id.txt"

def select_sites(stdscr, sites):
    logger.debug("Entering curses-based site selection menu with %d sites.", len(sites))
    curses.curs_set(0)
    current_row = 0
    selected = [False] * len(sites)
    top = 0
    while True:
        stdscr.clear()
        max_y, max_x = stdscr.getmaxyx()
        height = max_y
        for i in range(height):
            idx = top + i
            if idx >= len(sites):
                break
            prefix = "[x]" if selected[idx] else "[ ]"
            line = f"{prefix} {sites[idx]['name']}"[:max_x-1]
            if idx == current_row:
                stdscr.attron(curses.A_REVERSE)
            stdscr.addstr(i, 0, line)
            if idx == current_row:
                stdscr.attroff(curses.A_REVERSE)
        stdscr.refresh()
        key = stdscr.getch()
        if key == curses.KEY_UP and current_row > 0:
            current_row -= 1
            if current_row < top:
                top = current_row
        elif key == curses.KEY_DOWN and current_row < len(sites) - 1:
            current_row += 1
            if current_row >= top + height:
                top = current_row - height + 1
        elif key == ord(' '):
            selected[current_row] = not selected[current_row]
        elif key == ord('\n'):
            break
    selected_sites = [sites[i] for i in range(len(sites)) if selected[i]]
    logger.debug("Selected %d sites.", len(selected_sites))
    return selected_sites

def select_models(stdscr, models_list):
    logger.debug("Entering curses-based model selection menu with %d models.", len(models_list))
    curses.curs_set(0)
    current_row = 0
    selected = [False] * len(models_list)
    top = 0  # Scroll offset for models only
    while True:
        stdscr.clear()
        max_y, max_x = stdscr.getmaxyx()

        # Display prominent warning at the top
        stdscr.attron(curses.A_BOLD)
        stdscr.addstr(0, 0, "=== IMPORTANT WARNING ===")
        stdscr.addstr(1, 0, "ONLY select EITHER QFX or EX models.  Access and Core must not be done simultaneously. ")
        stdscr.addstr(2, 0, "Schedule at least 2hrs apart.  For sites with many access switches please use -b switch for access upgrade.")
        stdscr.attroff(curses.A_BOLD)
        stdscr.addstr(3, 0, "Use SPACE to toggle selection. Press ENTER when done.")
        stdscr.addstr(4, 0, "-" * min(max_x - 1, 60))

        # Calculate available rows for model list
        model_start_row = 5
        visible_height = max(1, max_y - model_start_row)

        # Display models
        for i in range(visible_height):
            idx = top + i
            if idx >= len(models_list):
                break
            prefix = "[x]" if selected[idx] else "[ ]"
            line = f"{prefix} {models_list[idx]}"[:max_x-1]
            display_row = model_start_row + i
            if idx == current_row:
                stdscr.attron(curses.A_REVERSE)
            stdscr.addstr(display_row, 0, line)
            if idx == current_row:
                stdscr.attroff(curses.A_REVERSE)

        stdscr.refresh()
        key = stdscr.getch()

        if key == curses.KEY_UP and current_row > 0:
            current_row -= 1
            if current_row < top:
                top = current_row
        elif key == curses.KEY_DOWN and current_row < len(models_list) - 1:
            current_row += 1
            if current_row >= top + visible_height:
                top = current_row - visible_height + 1
        elif key == ord(' '):
            selected[current_row] = not selected[current_row]
        elif key == ord('\n'):
            break

    selected_models = [models_list[i] for i in range(len(models_list)) if selected[i]]
    logger.debug("Selected %d models.", len(selected_models))
    return selected_models

# ... (rest of the script remains exactly the same as the previous version)

# Parse command-line arguments
parser = argparse.ArgumentParser(
    description="Juniper Mist Switch Upgrade Script\n\n"
                "This script allows upgrading switches across multiple sites in a Juniper Mist organization.\n"
                "It supports interactive site selection, model selection, upgrade initiation, status checking, and cancellation.\n\n"
                "Usage modes:\n"
                "  - Upgrade mode (default): Run without arguments to initiate an upgrade.\n"
                "    Prompts for API key, switch models (via menu), versions per model, start datetime, and site selection.\n"
                "  - Status mode: Use -s to check the status of the last upgrade and export to upgrade_status.xlsx.\n"
                "  - Cancel mode: Use -c to cancel the last upgrade.\n\n"
                "In upgrade mode, use -b for big_bang reboot strategy (default is serial). rrm is not supported for switches."
)
parser.add_argument('-s', '--status', action='store_true', help="Check upgrade status and export to upgrade_status.xlsx")
parser.add_argument('-c', '--cancel', action='store_true', help="Cancel the last upgrade")
parser.add_argument('-b', '--big-bang-reboot', action='store_true', help="Use big_bang reboot strategy (upgrade mode only)")

if '/?' in sys.argv:
    parser.print_help()
    sys.exit(0)

args = parser.parse_args()

logger.debug("Parsed arguments: %s", vars(args))

# Prompt for API key
api_key = getpass.getpass("Enter your API key: ")
logger.debug("API key entered (length: %d)", len(api_key))

# Set base URL and headers
base_url = "https://api.gc2.mist.com/api/v1"
headers = {
    "Authorization": f"Token {api_key}",
    "Content-Type": "application/json"
}
headers_no_content = {
    "Authorization": f"Token {api_key}"
}

# Fetch organization ID (shared across modes)
logger.debug("Fetching user details to retrieve organization ID.")
try:
    self_response = requests.get(f"{base_url}/self", headers=headers)
    self_response.raise_for_status()
    self_data = self_response.json()
    org_id = None
    for privilege in self_data.get("privileges", []):
        if privilege.get("scope") == "org":
            org_id = privilege.get("org_id")
            break
    if not org_id:
        logger.error("No organization ID found in user privileges.")
        sys.exit(1)
    logger.debug("Retrieved organization ID: %s", org_id)
except requests.exceptions.HTTPError as e:
    if e.response.status_code == 403:
        logger.error("Forbidden: Invalid API key or insufficient permissions.")
    else:
        logger.error("HTTP error fetching user details: %s", e)
    sys.exit(1)
except requests.exceptions.RequestException as e:
    logger.error("Error fetching user details: %s", e)
    sys.exit(1)
except Exception as e:
    logger.error("Unexpected error fetching user details: %s", e)
    sys.exit(1)

if args.status:
    logger.debug("Status mode activated.")
    if not os.path.exists(upgrade_id_file):
        logger.error("Upgrade ID file not found: %s", upgrade_id_file)
        sys.exit(1)
    try:
        with open(upgrade_id_file, "r") as f:
            upgrade_id = f.read().strip()
        logger.debug("Read upgrade ID: %s", upgrade_id)
    except Exception as e:
        logger.error("Error reading upgrade ID file: %s", e)
        sys.exit(1)

    # Fetch upgrade status
    try:
        status_response = requests.get(f"{base_url}/orgs/{org_id}/devices/upgrade/{upgrade_id}", headers=headers)
        status_response.raise_for_status()
        status_data = status_response.json()
        logger.info("Upgrade status:\n%s", json.dumps(status_data, indent=2))
    except requests.exceptions.HTTPError as e:
        if e.response.status_code == 403:
            logger.error("Forbidden: Insufficient permissions to fetch upgrade status.")
        else:
            logger.error("HTTP error fetching upgrade status: %s", e)
        sys.exit(1)
    except requests.exceptions.RequestException as e:
        logger.error("Error fetching upgrade status: %s", e)
        sys.exit(1)
    except Exception as e:
        logger.error("Unexpected error fetching upgrade status: %s", e)
        sys.exit(1)

    # Export to spreadsheet
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Upgrade Status"

        upgrades = status_data.get('upgrades', [])
        if not upgrades:
            logger.warning("No 'upgrades' key found in status data. Skipping export.")
        else:
            # Fetch sites for mapping
            try:
                sites_response = requests.get(f"{base_url}/orgs/{org_id}/sites", headers=headers)
                sites_response.raise_for_status()
                sites_data = sites_response.json()
                site_id_to_name = {s['id']: s['name'] for s in sites_data}
            except Exception as e:
                logger.error("Error fetching sites: %s", e)
                sys.exit(1)

            all_devices = []
            scheduled_time = status_data.get('start_datetime', 'N/A')
            reboot_time = status_data.get('reboot_datetime', 'N/A')
            download_strategy = status_data.get('download_strategy', 'N/A')
            reboot_strategy = status_data.get('reboot_strategy', 'N/A')
            for upgrade in upgrades:
                site_id = upgrade.get('site_id')
                site_name = site_id_to_name.get(site_id, 'Unknown')
                site_upgrade = upgrade.get('upgrade', {})
                devices = site_upgrade.get('targets', {}).get('devices', [])
                for dev in devices:
                    dev_copy = dev.copy()
                    dev_copy['site_name'] = site_name
                    dev_copy['site_id'] = site_id
                    dev_copy['site_upgrade_status'] = site_upgrade.get('status')
                    dev_copy['site_upgrade_id'] = site_upgrade.get('id')
                    dev_copy['scheduled_time'] = scheduled_time
                    dev_copy['reboot_time'] = reboot_time
                    dev_copy['download_strategy'] = download_strategy
                    dev_copy['reboot_strategy'] = reboot_strategy
                    all_devices.append(dev_copy)

            if not all_devices:
                logger.warning("No devices found in upgrades. Skipping export.")
            else:
                all_keys = set()
                for dev in all_devices:
                    all_keys.update(dev.keys())
                headers_list = sorted(list(all_keys))
                if 'site_name' in headers_list:
                    headers_list.remove('site_name')
                headers_list = ['site_name'] + headers_list

                for col, header in enumerate(headers_list, start=1):
                    ws.cell(row=1, column=col, value=header)

                row_num = 2
                for dev in all_devices:
                    for col, key in enumerate(headers_list, start=1):
                        value = dev.get(key)
                        ws.cell(row=row_num, column=col, value=str(value) if value is not None else '')
                    row_num += 1

                wb.save('upgrade_status.xlsx')
                logger.info("Exported upgrade status to upgrade_status.xlsx with %d rows.", row_num - 2)
    except Exception as e:
        logger.error("Error exporting to spreadsheet: %s", e)

elif args.cancel:
    logger.debug("Cancel mode activated.")
    if not os.path.exists(upgrade_id_file):
        logger.error("Upgrade ID file not found: %s", upgrade_id_file)
        sys.exit(1)
    try:
        with open(upgrade_id_file, "r") as f:
            upgrade_id = f.read().strip()
        logger.debug("Read upgrade ID: %s", upgrade_id)
    except Exception as e:
        logger.error("Error reading upgrade ID file: %s", e)
        sys.exit(1)

    # Fetch upgrade status to confirm
    try:
        status_response = requests.get(f"{base_url}/orgs/{org_id}/devices/upgrade/{upgrade_id}", headers=headers)
        status_response.raise_for_status()
        status_data = status_response.json()
        logger.debug("Fetched status data: %s", json.dumps(status_data, indent=2))
    except requests.exceptions.HTTPError as e:
        if e.response.status_code == 403:
            logger.error("Forbidden: Insufficient permissions to fetch upgrade status.")
        else:
            logger.error("HTTP error fetching upgrade status: %s", e)
        sys.exit(1)
    except requests.exceptions.RequestException as e:
        logger.error("Error fetching upgrade status: %s", e)
        sys.exit(1)
    except Exception as e:
        logger.error("Unexpected error fetching upgrade status: %s", e)
        sys.exit(1)

    # Cancel the upgrade
    try:
        cancel_url = f"{base_url}/orgs/{org_id}/devices/upgrade/{upgrade_id}/cancel"
        cancel_response = requests.post(cancel_url, headers=headers_no_content)
        cancel_response.raise_for_status()
        logger.info("Successfully canceled the upgrade.")
    except requests.exceptions.HTTPError as e:
        if e.response.status_code == 403:
            logger.error("Forbidden: Insufficient permissions to cancel upgrade.")
        else:
            logger.error("HTTP error canceling upgrade: %s", e)
        sys.exit(1)
    except requests.exceptions.RequestException as e:
        logger.error("Error canceling upgrade: %s", e)
        sys.exit(1)
    except Exception as e:
        logger.error("Unexpected error canceling upgrade: %s", e)
        sys.exit(1)

else:
    logger.debug("Upgrade mode activated.")

    # Fetch sites
    logger.debug("Fetching sites for organization %s.", org_id)
    try:
        sites_response = requests.get(f"{base_url}/orgs/{org_id}/sites", headers=headers)
        sites_response.raise_for_status()
        sites_data = sites_response.json()
        sites = [{'name': s['name'], 'id': s['id']} for s in sites_data]
        logger.debug("Fetched %d sites.", len(sites))
        if not sites:
            logger.error("No sites found in the organization.")
            sys.exit(1)
    except requests.exceptions.HTTPError as e:
        if e.response.status_code == 403:
            logger.error("Forbidden: Insufficient permissions to fetch sites.")
        else:
            logger.error("HTTP error fetching sites: %s", e)
        sys.exit(1)
    except requests.exceptions.RequestException as e:
        logger.error("Error fetching sites: %s", e)
        sys.exit(1)
    except Exception as e:
        logger.error("Unexpected error fetching sites: %s", e)
        sys.exit(1)

    # Select sites
    try:
        selected_sites = curses.wrapper(select_sites, sites)
        if not selected_sites:
            logger.error("No sites selected. Aborting upgrade.")
            sys.exit(1)
        site_ids = [s['id'] for s in selected_sites]
        logger.debug("Selected site IDs: %s", site_ids)
    except Exception as e:
        logger.error("Error during site selection: %s", e)
        sys.exit(1)

    # Fetch switch models from selected sites
    logger.debug("Fetching available switch models from selected sites.")
    switch_models_set = set()
    for site in selected_sites:
        site_id = site['id']
        try:
            devices_response = requests.get(f"{base_url}/sites/{site_id}/devices?type=switch", headers=headers)
            devices_response.raise_for_status()
            devices = devices_response.json()
            for dev in devices:
                if dev.get('type') == 'switch' and dev.get('model'):
                    switch_models_set.add(dev.get('model'))
        except Exception as e:
            logger.warning("Could not fetch devices for site %s: %s", site_id, e)
            continue

    if not switch_models_set:
        logger.error("No switch models found in the selected sites.")
        sys.exit(1)

    model_list = sorted(list(switch_models_set))
    logger.debug("Available switch models: %s", model_list)

    # Select models
    try:
        selected_models = curses.wrapper(select_models, model_list)
        if not selected_models:
            logger.error("No switch models selected. Aborting upgrade.")
            sys.exit(1)
        logger.debug("Selected switch models: %s", selected_models)
    except Exception as e:
        logger.error("Error during model selection: %s", e)
        sys.exit(1)

    # Prompt for version per selected model
    logger.debug("Prompting for firmware versions per model.")
    model_versions = {}
    for model in selected_models:
        ver = input(f"Enter Junos firmware version for {model}: ")
        model_versions[model] = ver.strip()
        logger.debug("Version for %s entered: %s", model, ver)
    versions = [
        {
            "firmware_type": "junos",
            "model_version": model_versions
        }
    ]

    # Prompt for start_datetime
    start_datetime = input("Enter start_datetime in ISO 8601 format (must include timezone and be at least 2 hours ahead): ")
    logger.debug("start_datetime entered: %s", start_datetime)

    # Validate start_datetime
    try:
        start_dt = datetime.fromisoformat(start_datetime)
        if start_dt.tzinfo is None:
            logger.error("start_datetime must include timezone (e.g., -07:00).")
            sys.exit(1)
        start_utc = start_dt.astimezone(timezone.utc)
        now_utc = datetime.now(timezone.utc)
        if start_utc < now_utc + timedelta(hours=2):
            logger.error("start_datetime must be at least 2 hours ahead of current time.")
            sys.exit(1)
        logger.debug("start_datetime validated successfully.")
    except ValueError as e:
        logger.error("Invalid start_datetime format: %s", e)
        sys.exit(1)
    except Exception as e:
        logger.error("Unexpected error validating start_datetime: %s", e)
        sys.exit(1)

    # Construct payload
    payload = {
        "device_type": "switch",
        "versions": versions,
        "all_sites": False,
        "site_ids": site_ids,
        "models": selected_models,
        "start_datetime": start_datetime,
        "download_strategy": "big_bang",
        "reboot_strategy": "serial"
    }
    if args.big_bang_reboot:
        payload["reboot_strategy"] = "big_bang"
        logger.debug("Set reboot_strategy to big_bang due to -b flag.")
    logger.debug("Constructed payload: %s", json.dumps(payload, indent=2))

    # Send upgrade request
    logger.debug("Sending upgrade request to API.")
    try:
        upgrade_response = requests.post(f"{base_url}/orgs/{org_id}/devices/upgrade", headers=headers, json=payload)
        upgrade_response.raise_for_status()
        upgrade_data = upgrade_response.json()
        upgrade_id = upgrade_data.get("id")
        if upgrade_id:
            logger.info("Upgrade initiated successfully. Upgrade ID: %s", upgrade_id)
            try:
                with open(upgrade_id_file, "w") as f:
                    f.write(upgrade_id)
                logger.debug("Saved upgrade ID to file: %s", upgrade_id_file)
            except Exception as e:
                logger.error("Error saving upgrade ID to file: %s", e)
        else:
            logger.warning("Upgrade response did not contain an ID: %s", upgrade_data)
    except requests.exceptions.HTTPError as e:
        if e.response.status_code == 403:
            logger.error("Forbidden: Insufficient permissions to initiate upgrade.")
        else:
            logger.error("HTTP error initiating upgrade: %s", e)
        sys.exit(1)
    except requests.exceptions.RequestException as e:
        logger.error("Error initiating upgrade: %s", e)
        sys.exit(1)
    except Exception as e:
        logger.error("Unexpected error initiating upgrade: %s", e)
        sys.exit(1)

logger.debug("Script execution completed.")